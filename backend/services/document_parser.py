"""
文档解析服务：将 .pdf / .docx / .doc / .txt / .md 文件解析为带结构信息的文本段列表。

每段格式：
    {"text": str, "page_no": int | None, "heading": str | None}

page_no 从 1 开始（PDF 有页码；DOCX / Markdown / TXT 无页码用 None）。
heading 为最近一级标题，无标题时为 None。

解析策略：
  - PDF: 先用 pypdf 提取文字，失败则降级到 OCR (RapidOCR)
  - DOCX: 用 python-docx 解析
  - DOC: 先用 LibreOffice 转换为 DOCX，再解析
  - TXT: 直接读取文本
  - MD: 按 Markdown 语法解析
"""
from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path
from typing import TypedDict

from core.logger import logger

__all__ = ["parse_document", "TextSegment"]


class TextSegment(TypedDict):
    text:    str
    page_no: int | None
    heading: str | None


# ------------------------------------------------------------------
# 公共入口
# ------------------------------------------------------------------

SUPPORTED_TYPES = {"pdf", "docx", "doc", "txt", "md", "markdown"}


def parse_document(file_path: str, file_type: str) -> list[TextSegment]:
    """
    根据文件类型分派解析逻辑。

    Args:
        file_path: 本地文件绝对路径
        file_type: 'pdf' | 'docx' | 'doc' | 'txt' | 'md'

    Returns:
        TextSegment 列表；每段代表一个有意义的文本块（段落/页面）。
    """
    ft = file_type.lower().lstrip(".")
    logger.info("Document parse started | file={path} | type={type}", path=file_path, type=ft)

    if ft == "pdf":
        segments = _parse_pdf(file_path)
    elif ft == "docx":
        segments = _parse_docx(file_path)
    elif ft == "doc":
        segments = _parse_doc(file_path)
    elif ft in ("txt", "text"):
        segments = _parse_txt(file_path)
    elif ft in ("md", "markdown"):
        segments = _parse_md(file_path)
    else:
        raise ValueError(f"不支持的文件类型: {file_type}，支持: {', '.join(sorted(SUPPORTED_TYPES))}")

    logger.info("Document parse finished | file={path} | segments={count}", path=file_path, count=len(segments))
    return segments


# ------------------------------------------------------------------
# PDF 解析（pypdf + OCR 降级）
# ------------------------------------------------------------------

_ocr_engine = None


def _get_ocr_engine():
    """延迟初始化 OCR 引擎（首次调用较慢）。"""
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
        logger.info("RapidOCR engine initialized")
    return _ocr_engine


def _parse_pdf_with_ocr(file_path: str) -> list[TextSegment]:
    """使用 OCR 解析扫描件 PDF。"""
    import numpy as np
    import pypdfium2 as pdfium

    ocr = _get_ocr_engine()
    pdf = pdfium.PdfDocument(file_path)
    segments: list[TextSegment] = []

    logger.info(f"OCR: processing {len(pdf)} pages from {file_path}")

    for page_idx in range(len(pdf)):
        page_no = page_idx + 1
        page = pdf[page_idx]
        # 2倍分辨率渲染，提高OCR准确率
        bitmap = page.render(scale=2)
        pil_image = bitmap.to_pil()
        img_array = np.array(pil_image)

        # OCR识别
        result, _ = ocr(img_array)

        if not result:
            continue

        # 合并同一页的文字块
        page_texts = []
        for line in result:
            text = line[1]  # result格式: [bbox, text, confidence]
            if text and text.strip():
                page_texts.append(text.strip())

        if page_texts:
            full_text = "\n".join(page_texts)
            segments.append({"text": full_text, "page_no": page_no, "heading": None})

    pdf.close()
    logger.info(f"OCR: extracted {len(segments)} segments")
    return segments


def _parse_pdf(file_path: str) -> list[TextSegment]:
    """解析 PDF 文件，先尝试 pypdf 提取文字，失败则降级到 OCR。"""
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    segments: list[TextSegment] = []
    total_chars = 0

    for page_no, page in enumerate(reader.pages, start=1):
        raw = page.extract_text() or ""
        total_chars += len(raw.strip())
        # 按自然段落分割（连续换行）
        paras = [p.strip() for p in re.split(r"\n{2,}", raw) if p.strip()]
        if not paras:
            text = raw.strip()
            if text:
                segments.append({"text": text, "page_no": page_no, "heading": None})
            continue
        for para in paras:
            segments.append({"text": para, "page_no": page_no, "heading": None})

    # 如果 pypdf 无法提取文字，尝试 OCR
    if len(reader.pages) > 0 and total_chars == 0:
        logger.info(f"PDF has no text layer, trying OCR: {file_path}")
        try:
            segments = _parse_pdf_with_ocr(file_path)
            if not segments:
                raise ValueError("OCR 也无法识别文字，可能是空白 PDF 或图片质量太差")
        except ImportError as e:
            raise ValueError(
                f"PDF 文件无法提取文字，可能是扫描件。OCR 依赖缺失: {e}"
            )
        except Exception as e:
            raise ValueError(f"OCR 识别失败: {e}")

    return segments


# ------------------------------------------------------------------
# DOCX 解析（python-docx）
# ------------------------------------------------------------------

# Word 内置样式中文/英文对照
_HEADING_STYLES = {
    "heading 1", "heading 2", "heading 3",
    "heading 4", "heading 5",
    "标题 1", "标题 2", "标题 3", "标题 4", "标题 5",
}


def _parse_docx(file_path: str) -> list[TextSegment]:
    """解析 .docx 文件。"""
    from docx import Document  # type: ignore

    doc = Document(file_path)
    segments: list[TextSegment] = []
    current_heading: str | None = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name.lower() if para.style else ""
        if style_name in _HEADING_STYLES:
            current_heading = text
            segments.append({"text": text, "page_no": None, "heading": None})
        else:
            segments.append({"text": text, "page_no": None, "heading": current_heading})

    return segments


# ------------------------------------------------------------------
# DOC 解析（LibreOffice 转换为 DOCX 后解析）
# ------------------------------------------------------------------

def _parse_doc(file_path: str) -> list[TextSegment]:
    """
    解析 .doc 文件（旧版 Word 97-2003 格式）。
    使用 LibreOffice 将 .doc 转换为 .docx，然后解析。
    """
    import shutil

    # 检查 LibreOffice 是否可用
    lo_cmd = None
    for cmd in ["libreoffice", "soffice", r"C:\Program Files\LibreOffice\program\soffice.exe"]:
        if shutil.which(cmd):
            lo_cmd = cmd
            break

    if lo_cmd is None:
        raise ValueError(
            "无法解析 .doc 格式：需要安装 LibreOffice。"
            "请将文件另存为 .docx 格式后重新上传，或安装 LibreOffice。"
        )

    # 使用 LibreOffice 转换为 docx
    with tempfile.TemporaryDirectory() as tmp_dir:
        try:
            result = subprocess.run(
                [lo_cmd, "--headless", "--convert-to", "docx", "--outdir", tmp_dir, file_path],
                capture_output=True,
                text=True,
                timeout=60,
            )
            if result.returncode != 0:
                raise ValueError(f"LibreOffice 转换失败: {result.stderr}")

            # 找到转换后的文件
            docx_files = list(Path(tmp_dir).glob("*.docx"))
            if not docx_files:
                raise ValueError("LibreOffice 转换后未找到 .docx 文件")

            segments = _parse_docx(str(docx_files[0]))
            logger.info(f"DOC converted and parsed: {len(segments)} segments")
            return segments

        except subprocess.TimeoutExpired:
            raise ValueError("LibreOffice 转换超时，请稍后重试或转为 .docx 格式上传")
        except FileNotFoundError:
            raise ValueError(
                "无法找到 LibreOffice，请将文件另存为 .docx 格式后重新上传"
            )


# ------------------------------------------------------------------
# TXT 解析
# ------------------------------------------------------------------

def _parse_txt(file_path: str) -> list[TextSegment]:
    """解析纯文本文件。"""
    # 尝试多种编码
    content = None
    for encoding in ["utf-8", "gbk", "gb2312", "gb18030", "latin-1"]:
        try:
            content = Path(file_path).read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue

    if content is None:
        raise ValueError("无法读取文本文件：编码识别失败")

    segments: list[TextSegment] = []

    # 按段落拆分（空行分隔）
    raw_paras = re.split(r"\n{2,}", content)

    for para in raw_paras:
        para = para.strip()
        if para:
            segments.append({"text": para, "page_no": None, "heading": None})

    # 如果没有段落（整个文件是一个段落），整体作为一个段落
    if not segments and content.strip():
        segments.append({"text": content.strip(), "page_no": None, "heading": None})

    return segments


# ------------------------------------------------------------------
# Markdown 解析
# ------------------------------------------------------------------

def _parse_md(file_path: str) -> list[TextSegment]:
    """解析 Markdown 文件。"""
    content = Path(file_path).read_text(encoding="utf-8")
    segments: list[TextSegment] = []
    current_heading: str | None = None

    # 按段落拆分（空行分隔）
    raw_paras = re.split(r"\n{2,}", content)

    for para in raw_paras:
        para = para.strip()
        if not para:
            continue
        # 检测 Markdown 标题（# / ## / ###）
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", para, re.MULTILINE)
        if heading_match:
            current_heading = heading_match.group(2).strip()
            segments.append({"text": para, "page_no": None, "heading": None})
        else:
            segments.append({"text": para, "page_no": None, "heading": current_heading})

    return segments
